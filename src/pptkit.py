"""小型版式工具：统一字体、颜色、页眉页脚，所有图形都是可编辑的 PowerPoint 形状。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

LATIN = "Helvetica Neue"
EA = "PingFang SC"

def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

INK = "141414"
BODY = "2B2B2B"
MUTED = "6B6B6B"
FAINT = "A0A0A0"
RULE = "D6D5CF"
PANEL = "F3F2EE"
ACCENT = "1F4AA8"
ACCENT_T = "DCE4F5"   # 浅蓝
ACCENT_M = "8EA6DA"   # 中蓝（部分具备）
GAP = "D9622B"        # 缺
GAP_T = "FBE7DD"
WHITE = "FFFFFF"
STATUS_FILL = {"done": ACCENT, "part": ACCENT_M, "gap": GAP}

SW, SH = 13.333, 7.5


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(WHITE)
    return s


def _font(run, size=None, color=None, bold=None, italic=None, latin=LATIN, ea=EA):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = rgb(color)
    f.name = latin
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
    e = etree.SubElement(rPr, qn("a:ea"))
    e.set("typeface", ea)


def text(slide, x, y, w, h, paras, size=12, color=BODY, bold=False, align="l",
         anchor="t", lsp=None, margin=0.0, space_after=0, wrap=True):
    """paras: str | list of paragraphs; a paragraph is str or list of runs (str | (str, dict))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras, size, color, bold, align, lsp, space_after)
    return tb


def fill_tf(tf, paras, size=12, color=BODY, bold=False, align="l", lsp=None, space_after=0):
    if isinstance(paras, str):
        paras = [paras]
    al = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = al
        if lsp:
            para.line_spacing = lsp
        para.space_after = Pt(space_after)
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            if isinstance(r, tuple):
                t, st = r
            else:
                t, st = r, {}
            run = para.add_run()
            run.text = t
            _font(run, st.get("size", size), st.get("color", color), st.get("bold", bold), st.get("italic"))


def box(slide, x, y, w, h, fill=None, line=None, lw=0.75, dash=False, shape="rect", radius=None):
    kind = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE,
            "oval": MSO_SHAPE.OVAL}[shape]
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if shape == "round":
        s.adjustments[0] = radius if radius is not None else 0.12
    style_shape(s, fill, line, lw, dash)
    s.text_frame.margin_left = s.text_frame.margin_right = Inches(0.06)
    s.text_frame.margin_top = s.text_frame.margin_bottom = Inches(0.03)
    return s


def style_shape(s, fill=None, line=None, lw=0.75, dash=False):
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    else:
        s.line.fill.background()


def label_box(slide, x, y, w, h, paras, fill=None, line=None, size=10, color=INK, bold=False,
              align="c", anchor="m", lw=0.75, dash=False, shape="rect"):
    s = box(slide, x, y, w, h, fill, line, lw, dash, shape)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, paras, size, color, bold, align)
    return s


def poly(slide, pts, fill=None, line=None, lw=0.75):
    """pts in inches."""
    ff = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    ff.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=True)
    s = ff.convert_to_shape()
    s.shadow.inherit = False
    style_shape(s, fill, line, lw)
    return s


def line(slide, x1, y1, x2, y2, color=RULE, lw=0.75, dash=False, arrow=False, head="triangle"):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", head)
        te.set("w", "med")
        te.set("len", "med")
    return c


def header(slide, section, title, deck="老项目运维提效 · 立项汇报"):
    text(slide, 0.6, 0.38, 6, 0.25, [[(section, {"bold": True})]], size=9.5, color=ACCENT)
    text(slide, 7.0, 0.38, 5.73, 0.25, deck, size=9, color=FAINT, align="r")
    text(slide, 0.6, 0.66, 12.1, 0.7, title, size=22, color=INK, bold=True)


def footer(slide, n, total, note=None):
    line(slide, 0.6, 7.0, 12.73, 7.0, RULE, 0.5)
    if note:
        text(slide, 0.6, 7.06, 11.2, 0.3, note, size=7.5, color=MUTED)
    text(slide, 11.9, 7.06, 0.83, 0.3, "%02d / %02d" % (n, total), size=8, color=FAINT, align="r")


def caption(slide, x, y, w, fig, txt):
    text(slide, x, y, w, 0.3, [[(fig + "  ", {"bold": True, "color": INK}), (txt, {})]], size=8.5, color=MUTED)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def glyph(slide, x, y, status, s=0.11):
    b = box(slide, x, y, s, s, STATUS_FILL[status], None)
    return b


def table(slide, x, y, w, col_w, rows, size=9, header_size=None, row_h=0.3, first_bold=False,
          colors=None, aligns=None):
    """rows[0] = header. Plain paper-style table: header rule + hairlines, no fills."""
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gt.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_h)
        for j in range(nc):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.06 if j else 0.0)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            v = rows[i][j]
            tf = cell.text_frame
            tf.word_wrap = True
            hdr = i == 0
            c = MUTED if hdr else (colors[j] if colors else BODY)
            b = hdr or (first_bold and j == 0)
            al = (aligns[j] if aligns else "l")
            # reset paragraph
            tf.paragraphs[0].text = ""
            fill_tf(tf, v if isinstance(v, list) else [v], (header_size or size) if hdr else size, c, b, al)
            tcPr = cell._tc.get_or_add_tcPr()
            for ch in list(tcPr):
                tcPr.remove(ch)
            for side in ("lnL", "lnR", "lnT", "lnB"):
                ln = etree.SubElement(tcPr, qn("a:" + side))
                if side == "lnB" or (side == "lnT" and i == 0):
                    heavy = (i == 0) or (i == nr - 1)
                    ln.set("w", str(Pt(1.0 if heavy else 0.5)))
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    clr = etree.SubElement(sf, qn("a:srgbClr"))
                    clr.set("val", INK if heavy else RULE)
                else:
                    ln.set("w", "0")
                    etree.SubElement(ln, qn("a:noFill"))
            etree.SubElement(tcPr, qn("a:noFill"))
    return gt


def strip_styles(prs):
    """去掉形状自带的主题样式（阴影、默认线色），所有颜色都已显式设置。"""
    for sl in prs.slides:
        for sp in sl.shapes:
            st = sp._element.find(qn("p:style"))
            if st is not None:
                sp._element.remove(st)


# ---- 第 3 页总起与后面各页的导航条：两套流程，各自编号
TRACK = [("运维", ["看板", "答复", "修复", "经过"]), ("重构", ["计划", "分批", "工单", "详情", "经过"])]
CHIP_W, CHIP_H, CHIP_G, GRP_LW = 0.62, 0.26, 0.05, 0.46


def chips(slide, x, y, gi, active=None, show_name=True):
    """画一组带编号的小标签；active 为当前页的序号（0 起），返回结束 x。"""
    name, items = TRACK[gi]
    if show_name:
        text(slide, x, y, GRP_LW, CHIP_H, name, size=10, color=ACCENT if active is not None else MUTED,
             bold=True, anchor="m")
        x += GRP_LW
    for i, t in enumerate(items):
        on = active == i
        label_box(slide, x, y, CHIP_W, CHIP_H, [[("%d " % (i + 1), {"bold": True}), (t, {})]],
                  fill=ACCENT if on else None, line=None if on else RULE, size=9,
                  color=WHITE if on else MUTED, shape="round")
        x += CHIP_W + CHIP_G
    return x


def tracker(slide, gi, idx, y=0.32):
    """页面右上角的导航：两组都画出来，当前页高亮。"""
    total = sum(GRP_LW + len(it) * (CHIP_W + CHIP_G) for _, it in TRACK) + 0.25
    x = 12.73 - total + CHIP_G
    for g in range(len(TRACK)):
        x = chips(slide, x, y, g, idx if g == gi else None) + 0.25
